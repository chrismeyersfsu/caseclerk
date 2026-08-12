from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from caseclerk_pipeline.extractors.docx import extract_docx
from caseclerk_pipeline.extractors.pdf import NeedsOcrError, extract_pdf
from caseclerk_pipeline.extractors.registry import get_extractor, is_unsupported
from caseclerk_pipeline.extractors.text import extract_text


def test_extract_docx_happy_path(tmp_path: Path) -> None:
    path = tmp_path / "letter.docx"
    doc = DocxDocument()
    doc.add_paragraph("The deposition on April 21, 2026 revealed a scheduling conflict.")
    doc.save(str(path))

    markdown = extract_docx(path)
    assert "April 21, 2026" in markdown
    assert "deposition" in markdown


def test_extract_pdf_happy_path(tmp_path: Path) -> None:
    path = tmp_path / "letter.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, text="The deposition on April 21, 2026 revealed a scheduling conflict.")
    pdf.output(str(path))

    text = extract_pdf(path)
    assert "deposition" in text.lower()
    assert "April 21, 2026" in text


def test_extract_pdf_without_text_layer_needs_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    pdf = FPDF()
    pdf.add_page()  # a page with no text, standing in for a scanned/image-only PDF
    pdf.output(str(path))

    with pytest.raises(NeedsOcrError):
        extract_pdf(path)


def test_extract_text_utf8(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Café notes re: deposition", encoding="utf-8")
    assert extract_text(path) == "Café notes re: deposition"


def test_extract_text_latin1_fallback(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_bytes("Café notes".encode("latin-1"))
    assert extract_text(path) == "Café notes"


def test_registry_maps_known_extensions() -> None:
    assert get_extractor(".docx") is extract_docx
    assert get_extractor(".pdf") is extract_pdf
    assert get_extractor(".txt") is extract_text
    assert get_extractor(".md") is extract_text
    assert get_extractor(".DOCX") is extract_docx
    assert get_extractor(".xyz") is None


def test_registry_flags_unsupported_extensions() -> None:
    assert is_unsupported(".msg")
    assert is_unsupported(".eml")
    assert is_unsupported(".rtf")
    assert is_unsupported(".xlsx")
    assert not is_unsupported(".docx")
    assert not is_unsupported(".pdf")
